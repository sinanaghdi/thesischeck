"""Best-effort DOCX parser retaining source IDs and OOXML format evidence."""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Iterator
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.parsing.normalization import normalize_for_comparison
from app.schemas.document import DocumentElement, NormalizedDocument, ParseDiagnostic

_CAPTION_RE = re.compile(
    r"^\s*(شکل|تصویر|جدول|الگوریتم|figure|fig\.?|table|algorithm)\s*[-:]?\s*"
    r"([\d۰-۹٠-٩]+(?:[-.][\d۰-۹٠-٩]+)*)",
    re.IGNORECASE,
)
_CHAPTER_RE = re.compile(
    r"^\s*(فصل\s+[\d۰-۹٠-٩]+|chapter\s+\d+|چکیده|abstract|منابع|مراجع|"
    r"references|فهرست\s+مطالب|contents)",
    re.IGNORECASE,
)


def _iter_block_items(document: Any) -> Iterator[Paragraph | Table]:
    """Yield body-level paragraphs and tables in their original order."""
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _font_metadata(paragraph: Paragraph) -> list[dict[str, Any]]:
    runs: list[dict[str, Any]] = []
    for run in paragraph.runs:
        rpr = run._r.rPr
        rfonts = rpr.rFonts if rpr is not None else None
        runs.append(
            {
                "text": run.text[:500],
                "font_name": run.font.name,
                "font_ascii": rfonts.get(qn("w:ascii")) if rfonts is not None else None,
                "font_hansi": rfonts.get(qn("w:hAnsi")) if rfonts is not None else None,
                "font_complex_script": rfonts.get(qn("w:cs")) if rfonts is not None else None,
                "font_east_asia": rfonts.get(qn("w:eastAsia")) if rfonts is not None else None,
                "font_size_pt": round(run.font.size.pt, 2) if run.font.size else None,
                "bold": run.bold,
                "italic": run.italic,
                "rtl": bool(run._r.xpath("./w:rPr/w:rtl")),
            }
        )
    return runs


def _section_metadata(document: Any) -> list[dict[str, Any]]:
    sections: list[dict[str, Any]] = []
    for index, section in enumerate(document.sections, start=1):
        header_paragraphs = [paragraph.text for paragraph in section.header.paragraphs]
        footer_paragraphs = [paragraph.text for paragraph in section.footer.paragraphs]
        header_instructions = section.header._element.xpath(".//w:instrText/text()")
        footer_instructions = section.footer._element.xpath(".//w:instrText/text()")
        header_fields = section.header._element.xpath(".//w:fldSimple/@w:instr")
        footer_fields = section.footer._element.xpath(".//w:fldSimple/@w:instr")
        sections.append(
            {
                "section": index,
                "page_width_cm": round(section.page_width.cm, 3) if section.page_width else None,
                "page_height_cm": round(section.page_height.cm, 3) if section.page_height else None,
                "top_margin_cm": round(section.top_margin.cm, 3) if section.top_margin else None,
                "bottom_margin_cm": (
                    round(section.bottom_margin.cm, 3) if section.bottom_margin else None
                ),
                "left_margin_cm": round(section.left_margin.cm, 3) if section.left_margin else None,
                "right_margin_cm": (
                    round(section.right_margin.cm, 3) if section.right_margin else None
                ),
                "header_distance_cm": (
                    round(section.header_distance.cm, 3) if section.header_distance else None
                ),
                "footer_distance_cm": (
                    round(section.footer_distance.cm, 3) if section.footer_distance else None
                ),
                "different_first_page_header_footer": section.different_first_page_header_footer,
                "header_text": " | ".join(header_paragraphs)[:1000],
                "footer_text": " | ".join(footer_paragraphs)[:1000],
                "header_has_page_field": any(
                    "PAGE" in instruction.upper()
                    for instruction in header_instructions + header_fields
                ),
                "footer_has_page_field": any(
                    "PAGE" in instruction.upper()
                    for instruction in footer_instructions + footer_fields
                ),
            }
        )
    return sections


def _ooxml_diagnostics(content: bytes) -> list[ParseDiagnostic]:
    diagnostics: list[ParseDiagnostic] = []
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
        if "word/footnotes.xml" in names:
            diagnostics.append(
                ParseDiagnostic(
                    code="footnotes_not_parsed",
                    message="وجود پانوشت تشخیص داده شد، اما متن پانوشت‌ها هنوز استخراج نشده است.",
                    severity="warning",
                )
            )
        if "word/endnotes.xml" in names:
            diagnostics.append(
                ParseDiagnostic(
                    code="endnotes_not_parsed",
                    message=(
                        "یادداشت‌های پایانی تشخیص داده شد، اما متن آن‌ها هنوز استخراج "
                        "نشده است."
                    ),
                    severity="warning",
                )
            )
        if "word/comments.xml" in names:
            diagnostics.append(
                ParseDiagnostic(
                    code="comments_not_parsed",
                    message="سند دارای نظرهای Word است؛ نظرها در این مرحله بررسی نمی‌شوند.",
                    severity="info",
                )
            )
    except (OSError, zipfile.BadZipFile):
        diagnostics.append(
            ParseDiagnostic(
                code="ooxml_metadata_incomplete",
                message="برخی اطلاعات داخلی سند قابل بررسی نبود.",
                severity="warning",
            )
        )
    return diagnostics


def parse_docx(content: bytes, filename: str) -> NormalizedDocument:
    """Parse a previously validated DOCX file without writing it to disk."""
    document = Document(io.BytesIO(content))
    diagnostics = _ooxml_diagnostics(content)
    elements: list[DocumentElement] = []
    heading_stack: list[str] = []
    source_index = 0
    heading_count = 0
    paragraph_count = 0
    table_count = 0
    figure_count = 0
    caption_count = 0
    has_toc_field = False
    has_equations = False

    for block in _iter_block_items(document):
        source_index += 1
        if isinstance(block, Table):
            table_count += 1
            row_texts = [
                " | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells)
                for row in block.rows
            ]
            table_text = "\n".join(row_texts).strip()
            column_count = max((len(row.cells) for row in block.rows), default=0)
            elements.append(
                DocumentElement(
                    element_id=f"t-{source_index:05d}",
                    kind="table",
                    text=table_text[:5000],
                    normalized_text=normalize_for_comparison(table_text)[:5000],
                    heading_path=heading_stack.copy(),
                    source_index=source_index,
                    metadata={"rows": len(block.rows), "columns": column_count},
                )
            )
            continue

        text = block.text.strip()
        has_graphic = bool(block._p.xpath(".//w:drawing | .//w:pict"))
        has_equation = bool(block._p.xpath(".//m:oMath | .//m:oMathPara"))
        if not text and not has_graphic and not has_equation:
            continue

        style_name = block.style.name if block.style is not None else "Normal"
        lower_style = style_name.lower()
        style_heading = lower_style.startswith("heading") or lower_style in {"title", "subtitle"}
        heading_level: int | None = None
        if style_heading:
            match = re.search(r"(\d+)", style_name)
            heading_level = int(match.group(1)) if match else (1 if lower_style == "title" else 2)
        elif _CHAPTER_RE.match(text):
            heading_level = 1

        kind = "heading" if heading_level else "paragraph"
        has_equations = has_equations or has_equation
        if has_equation and not text:
            kind = "equation"
        if has_graphic:
            figure_count += 1
            if not text:
                text = "[تصویر/شکل؛ محتوای گرافیکی استخراج نشده است]"
        caption_match = _CAPTION_RE.match(text) if text else None
        if caption_match:
            caption_count += 1
            if kind != "heading" and caption_match.group(1).lower() in {
                "شکل", "تصویر", "figure", "fig."
            }:
                kind = "figure"

        if kind == "heading":
            heading_count += 1
            level = max(1, min(heading_level or 1, 9))
            heading_stack = heading_stack[: level - 1]
            heading_stack.append(text)
        else:
            paragraph_count += 1

        paragraph_format = block.paragraph_format
        metadata: dict[str, Any] = {
            "style": style_name,
            "alignment": str(block.alignment) if block.alignment is not None else None,
            "keep_with_next": paragraph_format.keep_with_next,
            "page_break_before": paragraph_format.page_break_before,
            "line_spacing": (
                str(paragraph_format.line_spacing)
                if paragraph_format.line_spacing is not None
                else None
            ),
            "line_spacing_rule": (
                str(paragraph_format.line_spacing_rule)
                if paragraph_format.line_spacing_rule is not None
                else None
            ),
            "space_before_pt": (
                paragraph_format.space_before.pt if paragraph_format.space_before else None
            ),
            "space_after_pt": (
                paragraph_format.space_after.pt if paragraph_format.space_after else None
            ),
            "left_indent_pt": (
                paragraph_format.left_indent.pt if paragraph_format.left_indent else None
            ),
            "right_indent_pt": (
                paragraph_format.right_indent.pt if paragraph_format.right_indent else None
            ),
            "first_line_indent_pt": (
                paragraph_format.first_line_indent.pt
                if paragraph_format.first_line_indent
                else None
            ),
            "bidi": bool(block._p.xpath("./w:pPr/w:bidi")),
            "run_count": len(block.runs),
            "runs": _font_metadata(block),
            "has_drawing": has_graphic,
            "has_equation": has_equation,
            "is_caption": caption_match is not None,
            "caption_label": caption_match.group(1) if caption_match else None,
            "caption_number": caption_match.group(2) if caption_match else None,
            "has_toc_field": "TOC" in " ".join(block._p.xpath(".//w:instrText/text()")),
        }
        has_toc_field = has_toc_field or bool(metadata["has_toc_field"])
        elements.append(
            DocumentElement(
                element_id=f"p-{source_index:05d}",
                kind=kind,
                text=text,
                normalized_text=normalize_for_comparison(text),
                heading_level=heading_level,
                heading_path=heading_stack.copy(),
                page_number=None,
                source_index=source_index,
                metadata=metadata,
            )
        )

    if has_toc_field:
        diagnostics.append(
            ParseDiagnostic(
                code="toc_field_present",
                message=(
                    "فیلد فهرست مطالب تشخیص داده شد؛ صحت شماره صفحه‌ها از این داده "
                    "قابل تأیید نیست."
                ),
                severity="info",
            )
        )
    if has_equations:
        diagnostics.append(
            ParseDiagnostic(
                code="equations_partial",
                message=(
                    "یک یا چند ساختار معادله شناسایی شد؛ محتوای معادله ممکن است فقط "
                    "به‌صورت جای‌نگهدار نمایش داده شود."
                ),
                severity="warning",
            )
        )
    diagnostics.append(
        ParseDiagnostic(
            code="docx_pagination_unavailable",
            message=(
                "شماره صفحه در DOCX از روی فایل منبع به‌طور قابل اعتماد قابل محاسبه نیست؛ "
                "مکان‌ها با شناسه و مسیر عنوان گزارش می‌شوند."
            ),
            severity="info",
        )
    )

    return NormalizedDocument(
        source_format="docx",
        file_name=filename,
        elements=elements,
        diagnostics=diagnostics,
        page_count=None,
        parsing_complete=True,
        format_metadata={
            "sections": _section_metadata(document),
            "paragraph_count": paragraph_count,
            "heading_count": heading_count,
            "table_count": table_count,
            "figure_count": figure_count,
            "caption_count": caption_count,
            "has_toc_field": has_toc_field,
            "has_equations": has_equations,
            "styles_count": len(document.styles),
        },
    )
