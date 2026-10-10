from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app.parsing.dispatcher import parse_document

ROOT = Path(__file__).resolve().parents[2]
MAX_BYTES = 30 * 1024 * 1024


def minimal_pdf(text: str = "Hello ThesisCheck") -> bytes:
    """Build a one-page PDF fixture with computed xref offsets and no extra dependency."""
    safe_text = text.replace("\\", "\\\\").replace("(", r"\(").replace(")", r"\)")
    stream = f"BT /F1 18 Tf 50 700 Td ({safe_text}) Tj ET\n".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"endstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode())
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode())
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(
        (
            f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(pdf)


def test_parses_synthetic_docx_with_stable_heading_locations() -> None:
    content = (ROOT / "tests" / "fixtures" / "synthetic_bad.docx").read_bytes()
    parsed = parse_document("synthetic_bad.docx", content, max_size_bytes=MAX_BYTES)
    assert parsed.source_format == "docx"
    assert parsed.parsing_complete is True
    assert parsed.format_metadata["heading_count"] >= 3
    assert parsed.format_metadata["table_count"] == 0
    assert parsed.elements[0].element_id.startswith("p-")
    assert all(element.page_number is None for element in parsed.elements)
    assert any(item.code == "docx_pagination_unavailable" for item in parsed.diagnostics)
    assert any(
        "فصل اول" in heading
        for element in parsed.elements
        for heading in element.heading_path
    )


def test_docx_extracts_table_as_separate_element() -> None:
    source = Document()
    source.add_heading("Test chapter", level=1)
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "0.9"
    stream = BytesIO()
    source.save(stream)

    parsed = parse_document("table-test.docx", stream.getvalue(), max_size_bytes=MAX_BYTES)
    tables = [element for element in parsed.elements if element.kind == "table"]
    assert len(tables) == 1
    assert "Accuracy" in tables[0].text
    assert tables[0].metadata["rows"] == 2


def test_parses_text_pdf_with_page_locations() -> None:
    parsed = parse_document("sample.pdf", minimal_pdf(), max_size_bytes=MAX_BYTES)
    assert parsed.source_format == "pdf"
    assert parsed.page_count == 1
    assert parsed.elements
    assert parsed.elements[0].page_number == 1
    assert any(item.code == "pdf_style_checks_unavailable" for item in parsed.diagnostics)


def test_reports_no_text_layer_for_scanned_style_pdf() -> None:
    # A blank but structurally valid page is a simple stand-in for no selectable text.
    parsed = parse_document("blank.pdf", minimal_pdf(""), max_size_bytes=MAX_BYTES)
    assert any(item.code == "pdf_no_text_layer" for item in parsed.diagnostics)


def test_invalid_pdf_degrades_to_clear_error() -> None:
    with pytest.raises(ValueError, match="PDF"):
        parse_document("broken.pdf", b"%PDF-1.4\nnot a real PDF", max_size_bytes=MAX_BYTES)
