"""Create synthetic DOCX fixtures for future checker tests. Content is fictional."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "tests" / "fixtures"


def add_content(doc: Document, *, bad: bool) -> None:
    doc.add_heading("SYNTHETIC FIXTURE — NOT A REAL THESIS", level=0)
    doc.add_paragraph("This generated document is only for automated tests. Do not cite it.")
    doc.add_heading("چکیده", level=1)
    doc.add_paragraph(
        "هدف این نمونه بررسی ساختار متن است. روش نمونه بر پایه داده‌های ساختگی است "
        "و نتیجه‌ای واقعی ارائه نمی‌کند."
    )
    doc.add_paragraph("واژگان کلیدی: یادگیری ماشین، مجموعه‌داده، دقت")
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "This synthetic fixture describes a made-up procedure and contains no real "
        "research results."
    )
    doc.add_paragraph("Keywords: machine learning; dataset; accuracy")
    doc.add_heading("فصل اول: مقدمه", level=1)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bad:
        paragraph.add_run(
            "در این سند از يادگيري ماشین و یادگیری ماشین استفاده شده است. "
            "یک گزاره آزمایشی به منبع [1] اشاره می‌کند و گزاره دیگری به [3]."
        )
    else:
        paragraph.add_run(
            "در این سند از یادگیری ماشین به‌صورت یکدست استفاده شده است. "
            "ارجاع زیر تنها برای آزمون ساختاری و کاملاً ساختگی است: [1]."
        )
    doc.add_heading("فصل دوم: روش نمونه", level=1)
    doc.add_paragraph("این بخش ساختگی است و فقط برای آزمون استخراج عنوان‌ها ساخته شده است.")
    doc.add_paragraph("شکل 1-1 نمایشی ساختگی برای آزمون کپشن است.")
    doc.add_paragraph("این متن به شکل 1-1 اشاره می‌کند.")
    doc.add_heading("منابع آزمایشی — ساختگی", level=1)
    doc.add_paragraph("[1] SYNTHETIC TEST ENTRY — NOT A REAL PUBLICATION.")
    if bad:
        doc.add_paragraph("[1] SYNTHETIC DUPLICATE ENTRY — NOT A REAL PUBLICATION.")
    else:
        doc.add_paragraph("SYNTHETIC UNCITED ENTRY — NOT A REAL PUBLICATION.")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for filename, bad in (("synthetic_good.docx", False), ("synthetic_bad.docx", True)):
        document = Document()
        add_content(document, bad=bad)
        document.save(OUT / filename)
        print(f"Generated {OUT / filename}")


if __name__ == "__main__":
    main()
