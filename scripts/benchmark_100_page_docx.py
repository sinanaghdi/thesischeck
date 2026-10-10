"""Opt-in performance benchmark for parsing a synthetic 100-page DOCX.

Run from the repository root:
    python -m scripts.benchmark_100_page_docx

The document is generated in memory and is never written to disk. Its content is
fictional and intended only to measure the supervised parsing path.
"""

from __future__ import annotations

import argparse
import json
import time
from io import BytesIO

from docx import Document

from app.core.config import get_settings
from app.parsing.worker_client import parse_document_in_worker


def build_synthetic_docx(pages: int, paragraphs_per_page: int) -> bytes:
    """Build a compact DOCX with explicit page breaks and fictional paragraphs."""
    document = Document()
    document.core_properties.title = "Synthetic ThesisCheck performance fixture"
    document.core_properties.subject = "Not a real thesis or publication"

    for page_number in range(1, pages + 1):
        if page_number > 1:
            document.add_page_break()
        document.add_heading(f"Synthetic section {page_number:03d}", level=1)
        for paragraph_number in range(1, paragraphs_per_page + 1):
            document.add_paragraph(
                f"Fictional benchmark paragraph {page_number:03d}-{paragraph_number:02d}. "
                "This repeated sample text is used only to measure parsing latency and "
                "must not be interpreted as research or cited as a source."
            )

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pages", type=int, default=100, help="Logical pages to generate (1-500).")
    parser.add_argument(
        "--paragraphs-per-page",
        type=int,
        default=6,
        help="Fictional paragraphs on each logical page (1-100).",
    )
    parser.add_argument(
        "--max-seconds",
        type=float,
        default=30.0,
        help="Non-binding acceptance threshold for parse latency (default: 30 seconds).",
    )
    args = parser.parse_args()

    if not 1 <= args.pages <= 500:
        parser.error("--pages must be between 1 and 500")
    if not 1 <= args.paragraphs_per_page <= 100:
        parser.error("--paragraphs-per-page must be between 1 and 100")
    if args.max_seconds <= 0:
        parser.error("--max-seconds must be positive")

    settings = get_settings()
    max_size_bytes = settings.max_upload_mb * 1024 * 1024

    generation_start = time.perf_counter()
    payload = build_synthetic_docx(args.pages, args.paragraphs_per_page)
    generation_seconds = time.perf_counter() - generation_start
    if len(payload) > max_size_bytes:
        parser.error(
            f"Generated DOCX is {len(payload)} bytes, exceeding MAX_UPLOAD_MB "
            f"({settings.max_upload_mb} MiB). Reduce the workload or raise the configured limit."
        )

    parse_start = time.perf_counter()
    parsed = parse_document_in_worker(
        "synthetic-performance-benchmark.docx",
        payload,
        max_size_bytes=max_size_bytes,
        timeout_seconds=settings.parser_worker_timeout_seconds,
        cpu_seconds=settings.parser_worker_cpu_seconds,
        memory_mb=settings.parser_worker_memory_mb,
        concurrency=settings.parser_worker_concurrency,
    )
    parse_seconds = time.perf_counter() - parse_start

    results = {
        "benchmark": "supervised_docx_parse",
        "synthetic_only": True,
        "logical_pages_requested": args.pages,
        "paragraphs_per_page": args.paragraphs_per_page,
        "upload_bytes": len(payload),
        "generation_seconds": round(generation_seconds, 3),
        "parse_seconds": round(parse_seconds, 3),
        "element_count": len(parsed.elements),
        "diagnostic_count": len(parsed.diagnostics),
        "parser_reported_page_count": parsed.page_count,
        "acceptance_threshold_seconds": args.max_seconds,
        "threshold_passed": parse_seconds <= args.max_seconds,
    }
    print(json.dumps(results, ensure_ascii=False, indent=2))
    return 0 if results["threshold_passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
